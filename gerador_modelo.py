#!/usr/bin/env python3
"""
gerador_modelo.py
=================
Criação automatizada de modelos Word (.docx) e planilhas Excel (.xlsx) de exemplo.
Permite gerar modelos profissionais com bordas, cores, tipografia e tags Jinja2
para substituição posterior via docxtpl.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import List, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Paletas de cores para os temas de certificado
TEMAS = {
    "ouro_azul": {
        "primaria": RGBColor(26, 43, 76),       # Azul Marinho Profundo #1A2B4C
        "secundaria": RGBColor(197, 160, 89),   # Dourado Elegante #C5A059
        "texto": RGBColor(51, 51, 51),          # Grafite Escuro #333333
        "fundo_hex": "FDFBF7",                  # Off-white quente
        "borda_hex": "C5A059",
        "nome": "Ouro Real & Azul Marinho"
    },
    "esmeralda": {
        "primaria": RGBColor(15, 61, 46),       # Verde Esmeralda Escuro #0F3D2E
        "secundaria": RGBColor(176, 142, 60),   # Dourado Suave #B08E3C
        "texto": RGBColor(40, 40, 40),
        "fundo_hex": "F7FAF8",
        "borda_hex": "0F3D2E",
        "nome": "Verde Esmeralda Nobre"
    },
    "classico": {
        "primaria": RGBColor(30, 30, 30),       # Preto Clássico #1E1E1E
        "secundaria": RGBColor(120, 120, 120),  # Cinza Médio #787878
        "texto": RGBColor(50, 50, 50),
        "fundo_hex": "FFFFFF",
        "borda_hex": "1E1E1E",
        "nome": "Clássico Minimalista"
    },
    "rubi": {
        "primaria": RGBColor(107, 24, 43),      # Borgonha/Rubi #6B182B
        "secundaria": RGBColor(212, 175, 55),   # Ouro Metálico #D4AF37
        "texto": RGBColor(45, 45, 45),
        "fundo_hex": "FCF8F9",
        "borda_hex": "6B182B",
        "nome": "Rubi Imperial"
    }
}


def _adicionar_borda_pagina(section, cor_hex: str = "C5A059"):
    """Adiciona borda dupla decorativa em torno da página do Word via XML."""
    try:
        sectPr = section._sectPr
        # Borda dupla ornamentada na página
        pgBorders_xml = f"""
        <w:pgBorders {nsdecls('w')} w:offsetFrom="page">
            <w:top w:val="double" w:sz="24" w:space="24" w:color="{cor_hex}"/>
            <w:left w:val="double" w:sz="24" w:space="24" w:color="{cor_hex}"/>
            <w:bottom w:val="double" w:sz="24" w:space="24" w:color="{cor_hex}"/>
            <w:right w:val="double" w:sz="24" w:space="24" w:color="{cor_hex}"/>
        </w:pgBorders>
        """
        sectPr.append(parse_xml(pgBorders_xml))
    except Exception as e:
        # Se der erro ao inserir borda de página, não impede o documento de ser criado
        pass


def _remover_bordas_tabela(table):
    """Remove todas as bordas visíveis de uma tabela Word via XML."""
    try:
        tblPr = table._tbl.tblPr
        borders_xml = f"""
        <w:tblBorders {nsdecls('w')}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
        """
        tblPr.append(parse_xml(borders_xml))
    except Exception:
        pass


def criar_modelo_docx(
    caminho_saida: Path | str,
    titulo: str = "CERTIFICADO",
    subtitulo: str = "DE CONCLUSÃO E PARTICIPAÇÃO",
    instituicao: str = "CENTRO DE DESENVOLVIMENTO & CAPACITAÇÃO PROFISSIONAL",
    texto_corpo: Optional[str] = None,
    assinatura_1_nome: str = "Prof. Coordenador(a)",
    assinatura_1_cargo: str = "Coordenação Pedagógica",
    assinatura_2_nome: str = "Diretoria Acadêmica",
    assinatura_2_cargo: str = "Direção Geral",
    tema: str = "ouro_azul",
    orientacao: str = "paisagem",
    codigo_verificacao: bool = True,
    logo_prestadora_path: Optional[str] = None,
    logo_contratante_path: Optional[str] = None,
    habilitar_verso: bool = False,
    verso_titulo: str = "CONTEÚDO PROGRAMÁTICO & REGISTRO",
    verso_conteudo: Optional[str] = None,
    verso_livro: str = "01",
    verso_folha: str = "45",
    verso_registro: str = "REG-{{CPF}}",
    verso_amparo_legal: Optional[str] = None,
    verso_assinatura_nome: str = "Coordenação Técnica Pedagógica",
    verso_assinatura_cargo: str = "Responsável Técnico(a)"
) -> Path:
    """
    Cria um documento Word (.docx) elegante para servir de modelo com Jinja2 tags,
    com suporte a logotipos da prestadora e contratante na frente e verso (dupla face).
    """
    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    paleta = TEMAS.get(tema, TEMAS["ouro_azul"])

    if not texto_corpo:
        texto_corpo = (
            "Certificamos para os devidos fins que {{NOME}}, inscrito(a) sob o CPF nº {{CPF}}, "
            "concluiu com êxito o curso de {{CURSO}}, realizado em {{DATA}}, "
            "com aproveitamento satisfatório e carga horária total de {{CARGA_HORARIA}}."
        )

    if not verso_conteudo:
        verso_conteudo = (
            "• Módulo 1: Introdução, Fundamentos e Conceitos Básicos\n"
            "• Módulo 2: Metodologias Ágeis, Ferramentas e Boas Práticas\n"
            "• Módulo 3: Estudos de Casos Práticos e Aplicações no Mercado\n"
            "• Módulo 4: Avaliação Final, Segurança e Ética Profissional"
        )

    if not verso_amparo_legal:
        verso_amparo_legal = (
            "Certificado emitido e registrado em estrita conformidade com a Lei Federal de "
            "Diretrizes e Bases da Educação Nacional (Lei nº 9.394/96) e Decreto Presidencial nº 5.154/04."
        )

    doc = docx.Document()
    section = doc.sections[0]

    # Orientação e Dimensões (A4 Paisagem padrão)
    is_paisagem = orientacao.lower() == "paisagem"
    if is_paisagem:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        largura_util = 9.89
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        largura_util = 6.67

    # Borda na página
    _adicionar_borda_pagina(section, paleta["borda_hex"])

    # =========================================================================
    # FRENTE DO CERTIFICADO
    # =========================================================================

    # 1. Cabeçalho com Logotipos (Prestadora à Esquerda, Contratante à Direita)
    tem_logo_esq = logo_prestadora_path and Path(logo_prestadora_path).is_file()
    tem_logo_dir = logo_contratante_path and Path(logo_contratante_path).is_file()

    if tem_logo_esq or tem_logo_dir:
        tab_logos = doc.add_table(rows=1, cols=3)
        tab_logos.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab_logos.autofit = False
        _remover_bordas_tabela(tab_logos)

        larg_logo = Inches(2.2 if is_paisagem else 1.6)
        larg_centro = Inches(largura_util - (2.2 * 2 if is_paisagem else 1.6 * 2))

        for row in tab_logos.rows:
            row.cells[0].width = larg_logo
            row.cells[1].width = larg_centro
            row.cells[2].width = larg_logo
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Coluna 0: Logo Prestadora (Esquerda)
        p0 = tab_logos.cell(0, 0).paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        if tem_logo_esq:
            try:
                p0.add_run().add_picture(str(logo_prestadora_path), height=Inches(0.75))
            except Exception:
                pass

        # Coluna 1: Nome da Instituição (Centro)
        p1 = tab_logos.cell(0, 1).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r_inst = p1.add_run(instituicao.upper())
        r_inst.font.name = "Georgia"
        r_inst.font.size = Pt(10.5)
        r_inst.font.bold = True
        r_inst.font.color.rgb = paleta["secundaria"]

        # Coluna 2: Logo Contratante (Direita)
        p2 = tab_logos.cell(0, 2).paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        if tem_logo_dir:
            try:
                p2.add_run().add_picture(str(logo_contratante_path), height=Inches(0.75))
            except Exception:
                pass
    else:
        # Sem logotipos: cabeçalho clássico centralizado
        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_inst.paragraph_format.space_before = Pt(10)
        p_inst.paragraph_format.space_after = Pt(6)
        run_inst = p_inst.add_run(instituicao.upper())
        run_inst.font.name = "Georgia"
        run_inst.font.size = Pt(11)
        run_inst.font.bold = True
        run_inst.font.color.rgb = paleta["secundaria"]

    # 2. Título Principal
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(4)
    p_tit.paragraph_format.space_after = Pt(2)
    run_tit = p_tit.add_run(titulo.upper())
    run_tit.font.name = "Georgia"
    run_tit.font.size = Pt(30 if is_paisagem else 26)
    run_tit.font.bold = True
    run_tit.font.color.rgb = paleta["primaria"]

    # 3. Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run(subtitulo.upper())
    run_sub.font.name = "Georgia"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = False
    run_sub.font.color.rgb = paleta["secundaria"]

    # 4. Linha divisória ornamental
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(18)
    run_div = p_div.add_run("♦   ❖   ♦")
    run_div.font.name = "Georgia"
    run_div.font.size = Pt(11)
    run_div.font.color.rgb = paleta["secundaria"]

    # 5. Corpo do Texto com quebras e formatação inteligente
    p_corpo = doc.add_paragraph()
    p_corpo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_corpo.paragraph_format.line_spacing = 1.35
    p_corpo.paragraph_format.space_before = Pt(0)
    p_corpo.paragraph_format.space_after = Pt(60 if not is_paisagem else 36)

    import re
    partes = re.split(r"(\{\{[A-Za-z0-9_]+\}\})", texto_corpo)
    for parte in partes:
        if not parte:
            continue
        run = p_corpo.add_run(parte)
        run.font.name = "Calibri"
        run.font.size = Pt(13.5)
        if parte.startswith("{{") and parte.endswith("}}"):
            run.font.bold = True
            run.font.color.rgb = paleta["primaria"]
        else:
            run.font.bold = False
            run.font.color.rgb = paleta["texto"]

    # 6. Tabela de Assinaturas (2 colunas alinhadas)
    tabela = doc.add_table(rows=2, cols=2)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    _remover_bordas_tabela(tabela)

    largura_col = Inches((largura_util - 0.5) / 2)
    for row in tabela.rows:
        for cell in row.cells:
            cell.width = largura_col
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Linha 0: Traços de assinatura
    c0 = tabela.cell(0, 0).paragraphs[0]
    c0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c0.paragraph_format.space_after = Pt(4)
    r_traco1 = c0.add_run("_________________________________________")
    r_traco1.font.color.rgb = paleta["secundaria"]

    c1 = tabela.cell(0, 1).paragraphs[0]
    c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c1.paragraph_format.space_after = Pt(4)
    r_traco2 = c1.add_run("_________________________________________")
    r_traco2.font.color.rgb = paleta["secundaria"]

    # Linha 1: Nomes e Cargos
    c2 = tabela.cell(1, 0).paragraphs[0]
    c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c2.paragraph_format.space_after = Pt(0)
    r_n1 = c2.add_run(f"{assinatura_1_nome}\n")
    r_n1.font.name = "Georgia"
    r_n1.font.size = Pt(10.5)
    r_n1.font.bold = True
    r_n1.font.color.rgb = paleta["primaria"]
    r_c1 = c2.add_run(assinatura_1_cargo)
    r_c1.font.name = "Calibri"
    r_c1.font.size = Pt(9)
    r_c1.font.color.rgb = paleta["texto"]

    c3 = tabela.cell(1, 1).paragraphs[0]
    c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c3.paragraph_format.space_after = Pt(0)
    r_n2 = c3.add_run(f"{assinatura_2_nome}\n")
    r_n2.font.name = "Georgia"
    r_n2.font.size = Pt(10.5)
    r_n2.font.bold = True
    r_n2.font.color.rgb = paleta["primaria"]
    r_c2 = c3.add_run(assinatura_2_cargo)
    r_c2.font.name = "Calibri"
    r_c2.font.size = Pt(9)
    r_c2.font.color.rgb = paleta["texto"]

    # 7. Rodapé com autenticação
    if codigo_verificacao:
        p_rodape = doc.add_paragraph()
        p_rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_rodape.paragraph_format.space_before = Pt(20)
        p_rodape.paragraph_format.space_after = Pt(0)
        r_rod = p_rodape.add_run("Registro de Certificado Digital • Válido em todo território nacional")
        r_rod.font.name = "Calibri"
        r_rod.font.size = Pt(8.5)
        r_rod.font.italic = True
        r_rod.font.color.rgb = RGBColor(140, 140, 140)

    # =========================================================================
    # VERSO DO CERTIFICADO (FOLHA DE TRÁS / SEGUNDA PÁGINA)
    # =========================================================================
    if habilitar_verso:
        doc.add_page_break()

        # Cabeçalho do Verso
        p_v_inst = doc.add_paragraph()
        p_v_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v_inst.paragraph_format.space_before = Pt(10)
        p_v_inst.paragraph_format.space_after = Pt(4)
        r_vi = p_v_inst.add_run(instituicao.upper())
        r_vi.font.name = "Georgia"
        r_vi.font.size = Pt(10)
        r_vi.font.bold = True
        r_vi.font.color.rgb = paleta["secundaria"]

        p_v_tit = doc.add_paragraph()
        p_v_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v_tit.paragraph_format.space_before = Pt(0)
        p_v_tit.paragraph_format.space_after = Pt(6)
        r_vt = p_v_tit.add_run(verso_titulo.upper())
        r_vt.font.name = "Georgia"
        r_vt.font.size = Pt(18 if is_paisagem else 16)
        r_vt.font.bold = True
        r_vt.font.color.rgb = paleta["primaria"]

        p_v_div = doc.add_paragraph()
        p_v_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v_div.paragraph_format.space_before = Pt(0)
        p_v_div.paragraph_format.space_after = Pt(18)
        r_vd = p_v_div.add_run("♦   ❖   ♦")
        r_vd.font.name = "Georgia"
        r_vd.font.size = Pt(10)
        r_vd.font.color.rgb = paleta["secundaria"]

        # Tabela em 2 colunas para o verso
        tab_verso = doc.add_table(rows=1, cols=2)
        tab_verso.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab_verso.autofit = False
        _remover_bordas_tabela(tab_verso)

        larg_col_esq = Inches(5.6 if is_paisagem else 3.8)
        larg_col_dir = Inches(largura_util - (5.6 if is_paisagem else 3.8) - 0.3)

        for row in tab_verso.rows:
            row.cells[0].width = larg_col_esq
            row.cells[1].width = larg_col_dir
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

        # Coluna 0: Conteúdo Programático / Ementa
        c_esq = tab_verso.cell(0, 0)
        p_e_tit = c_esq.paragraphs[0]
        p_e_tit.paragraph_format.space_after = Pt(8)
        r_et = p_e_tit.add_run("CONTEÚDO PROGRAMÁTICO MINISTRADO:")
        r_et.font.name = "Georgia"
        r_et.font.size = Pt(11)
        r_et.font.bold = True
        r_et.font.color.rgb = paleta["primaria"]

        # Divide as linhas da ementa
        linhas_ementa = [l.strip() for l in verso_conteudo.strip().split("\n") if l.strip()]
        for linha in linhas_ementa:
            p_mod = c_esq.add_paragraph()
            p_mod.paragraph_format.space_before = Pt(0)
            p_mod.paragraph_format.space_after = Pt(5)
            p_mod.paragraph_format.line_spacing = 1.2
            
            # Adiciona com substituição de tags se houver
            partes_mod = re.split(r"(\{\{[A-Za-z0-9_]+\}\})", linha)
            for pm in partes_mod:
                if not pm:
                    continue
                rm = p_mod.add_run(pm)
                rm.font.name = "Calibri"
                rm.font.size = Pt(10)
                if pm.startswith("{{") and pm.endswith("}}"):
                    rm.font.bold = True
                    rm.font.color.rgb = paleta["primaria"]
                else:
                    rm.font.color.rgb = paleta["texto"]

        # Resumo de Carga Horária na coluna da esquerda
        p_ch = c_esq.add_paragraph()
        p_ch.paragraph_format.space_before = Pt(12)
        p_ch.paragraph_format.space_after = Pt(0)
        r_ch_lbl = p_ch.add_run("Carga Horária Total Integralizada: ")
        r_ch_lbl.font.name = "Calibri"
        r_ch_lbl.font.size = Pt(10)
        r_ch_lbl.font.bold = True
        r_ch_lbl.font.color.rgb = paleta["primaria"]
        r_ch_val = p_ch.add_run("{{CARGA_HORARIA}}")
        r_ch_val.font.name = "Calibri"
        r_ch_val.font.size = Pt(10)
        r_ch_val.font.bold = True
        r_ch_val.font.color.rgb = paleta["secundaria"]

        # Coluna 1: Dados de Registro, Amparo Legal e Assinatura Técnica
        c_dir = tab_verso.cell(0, 1)
        p_r_tit = c_dir.paragraphs[0]
        p_r_tit.paragraph_format.space_after = Pt(8)
        r_rt = p_r_tit.add_run("REGISTRO & CONFORMIDADE:")
        r_rt.font.name = "Georgia"
        r_rt.font.size = Pt(11)
        r_rt.font.bold = True
        r_rt.font.color.rgb = paleta["primaria"]

        # Box de Registro
        p_box_reg = c_dir.add_paragraph()
        p_box_reg.paragraph_format.space_after = Pt(10)
        p_box_reg.paragraph_format.line_spacing = 1.3
        
        reg_itens = [
            ("Livro de Registro: ", verso_livro),
            ("Folha: ", verso_folha),
            ("Nº de Registro: ", verso_registro),
            ("Data de Emissão: ", "{{DATA}}")
        ]
        for rotulo, val in reg_itens:
            r_l = p_box_reg.add_run(rotulo)
            r_l.font.name = "Calibri"
            r_l.font.size = Pt(9.5)
            r_l.font.bold = True
            r_l.font.color.rgb = paleta["texto"]
            
            partes_val = re.split(r"(\{\{[A-Za-z0-9_]+\}\})", val)
            for pv in partes_val:
                if not pv:
                    continue
                r_v = p_box_reg.add_run(pv)
                r_v.font.name = "Calibri"
                r_v.font.size = Pt(9.5)
                if pv.startswith("{{") and pv.endswith("}}"):
                    r_v.font.bold = True
                    r_v.font.color.rgb = paleta["primaria"]
                else:
                    r_v.font.color.rgb = paleta["texto"]
            p_box_reg.add_run("\n")

        # Amparo Legal
        p_amp = c_dir.add_paragraph()
        p_amp.paragraph_format.space_before = Pt(4)
        p_amp.paragraph_format.space_after = Pt(16)
        p_amp.paragraph_format.line_spacing = 1.2
        r_amp = p_amp.add_run(verso_amparo_legal)
        r_amp.font.name = "Calibri"
        r_amp.font.size = Pt(8.5)
        r_amp.font.italic = True
        r_amp.font.color.rgb = RGBColor(100, 100, 100)

        # Assinatura Técnica do Verso
        p_v_ass = c_dir.add_paragraph()
        p_v_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v_ass.paragraph_format.space_before = Pt(10)
        p_v_ass.paragraph_format.space_after = Pt(0)
        r_vt_traco = p_v_ass.add_run("_____________________________________\n")
        r_vt_traco.font.color.rgb = paleta["secundaria"]
        r_vt_nome = p_v_ass.add_run(f"{verso_assinatura_nome}\n")
        r_vt_nome.font.name = "Georgia"
        r_vt_nome.font.size = Pt(10)
        r_vt_nome.font.bold = True
        r_vt_nome.font.color.rgb = paleta["primaria"]
        r_vt_cargo = p_v_ass.add_run(verso_assinatura_cargo)
        r_vt_cargo.font.name = "Calibri"
        r_vt_cargo.font.size = Pt(8.5)
        r_vt_cargo.font.color.rgb = paleta["texto"]

    doc.save(str(caminho_saida))
    return caminho_saida


def criar_planilha_exemplo(
    caminho_saida: Path | str,
    colunas: Optional[List[str]] = None
) -> Path:
    """
    Gera uma planilha Excel estilizada e com dados de teste realistas.
    """
    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    if not colunas:
        colunas = ["NOME", "CPF", "CURSO", "DATA", "CARGA_HORARIA"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Participantes"

    # Cabeçalho estilizado
    fill_cabecalho = PatternFill(start_color="1A2B4C", end_color="1A2B4C", fill_type="solid")
    font_cabecalho = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_centro = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_esquerda = Alignment(horizontal="left", vertical="center")
    borda_fina = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    ws.append(colunas)

    for col_idx in range(1, len(colunas) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = fill_cabecalho
        cell.font = font_cabecalho
        cell.alignment = align_centro
        cell.border = borda_fina

    # Dados de exemplo pré-configurados
    exemplos = [
        {
            "NOME": "Ana Beatriz da Silva Santos",
            "CPF": "123.456.789-01",
            "CURSO": "Inteligência Artificial & Automação",
            "DATA": "11/09/2026",
            "CARGA_HORARIA": "40 horas"
        },
        {
            "NOME": "Carlos Eduardo de Souza Mendes",
            "CPF": "234.567.890-12",
            "CURSO": "Gestão Ágil e Liderança de Equipes",
            "DATA": "11/09/2026",
            "CARGA_HORARIA": "32 horas"
        },
        {
            "NOME": "Mariana Oliveira Albuquerque",
            "CPF": "345.678.901-23",
            "CURSO": "Segurança da Informação & LGPD",
            "DATA": "11/09/2026",
            "CARGA_HORARIA": "20 horas"
        }
    ]

    fill_zebrado = PatternFill(start_color="F7F9FB", end_color="F7F9FB", fill_type="solid")

    for row_idx, dado in enumerate(exemplos, start=2):
        linha = []
        for col in colunas:
            col_upper = col.strip().upper()
            linha.append(dado.get(col_upper, f"Valor de {col}"))
        ws.append(linha)

        for col_idx in range(1, len(colunas) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.alignment = align_esquerda if col_idx == 1 else align_centro
            cell.border = borda_fina
            if row_idx % 2 == 1:
                cell.fill = fill_zebrado

    # Ajusta largura automática das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    ws.row_dimensions[1].height = 26
    for r in range(2, len(exemplos) + 2):
        ws.row_dimensions[r].height = 22

    wb.save(str(caminho_saida))
    return caminho_saida


if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    mod_path = base / "modelo" / "certificado_base.docx"
    xls_path = base / "entrada" / "participantes_exemplo.xlsx"
    print(f"Gerando modelo padrão em: {mod_path}")
    criar_modelo_docx(mod_path)
    print(f"Gerando planilha de exemplo em: {xls_path}")
    criar_planilha_exemplo(xls_path)
    print("Concluído!")
